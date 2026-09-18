from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("reports") / "MailSense_Academic_Project_Report.docx"
BLACK = RGBColor(0, 0, 0)
GRAY = "808080"


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill="FFFFFF"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="000000", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_para(doc, text="", style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, 11, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    for cell, text in zip(hdr.cells, headers):
        set_cell_shading(cell, "E6E6E6")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, 10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_shading(cell, "FFFFFF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, 10)
    # Repeat header row across pages.
    tr_pr = hdr._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    doc.add_paragraph()
    return table


def add_heading(doc, number, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(f"{number} {text}")
    set_run_font(r, 14 if level == 1 else 12, bold=True)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.49)
section.footer_distance = Inches(0.49)

# Formal academic black-and-white token set, derived from the restrained preset.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(11)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for name, size, before, after in (("Heading 1", 14, 14, 6), ("Heading 2", 12, 10, 4), ("Heading 3", 11, 8, 3)):
    s = styles[name]
    s.font.name = "Times New Roman"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = BLACK
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
for list_name in ("List Bullet", "List Number"):
    s = styles[list_name]
    s.font.name = "Times New Roman"
    s.font.size = Pt(11)
    s.font.color.rgb = BLACK
    s.paragraph_format.space_after = Pt(3)

footer = section.footer.paragraphs[0]
add_page_number(footer)

# Title page
for _ in range(7):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MailSense: Actionable Email Detection Using NLP")
set_run_font(r, 20, bold=True)
p.paragraph_format.space_after = Pt(18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Formal Academic Project Report")
set_run_font(r, 13)
p.paragraph_format.space_after = Pt(40)
for line in ("Md Ariful Alam Mahim (Roll: 2107023)", "Md Jubair Husain (Roll: 2107029)"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(line)
    set_run_font(r, 12)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(46)
r = p.add_run("September 2026")
set_run_font(r, 11)
doc.add_section(WD_SECTION.NEW_PAGE)

add_heading(doc, "1.", "Introduction")
add_para(doc, "Electronic mail can contain requests, deadlines, confirmations, notices, and other content that may require a recipient's attention. MailSense addresses the focused task of identifying whether a piece of email content is actionable. The project implements and compares three binary text-classification pipelines: a traditional TF-IDF plus Logistic Regression baseline, a pretrained Word2Vec plus bidirectional LSTM (BiLSTM) neural model, and a fine-tuned pretrained BERT model (bert-base-uncased).")
add_para(doc, "The repository describes the input records as individual email sentences or snippets rather than complete messages with separate subject and body fields. Accordingly, this report uses the precise term actionable email-content classification. The original labels are retained: Yes denotes Actionable and No denotes Non-Actionable.")

add_heading(doc, "2.", "Problem Statement")
add_para(doc, "Given a text snippet from an email dataset, the system must assign one of two labels: Actionable (Yes) or Non-Actionable (No). The positive class is Actionable (Yes). The experimental question is how the three implemented approaches perform when trained and evaluated using the same cleaned data split and held-out test set.")

add_heading(doc, "3.", "Objectives")
for item in (
    "Prepare a cleaned, labeled dataset and create common stratified train, validation, and test splits.",
    "Implement three models representing sparse statistical features, pretrained word embeddings with sequential modeling, and pretrained Transformer fine-tuning.",
    "Evaluate every model on the same held-out test split using accuracy, precision, recall, F1-score, ROC-AUC, and confusion-matrix counts.",
    "Compare the final test results while treating Actionable (Yes) as the positive class.",
):
    add_bullet(doc, item)

add_heading(doc, "4.", "Dataset")
add_para(doc, "The source dataset is the Parakweet Labs Email Intent Dataset file Ask0729-fixed.txt. In the original project dataset, there are 3,657 examples: 1,719 Actionable (Yes) and 1,938 Non-Actionable (No). These figures describe the source data before the project cleaning pipeline is applied.")
add_para(doc, "For the project experiments, the preprocessing code produced 3,651 final examples after removing six records. The retained experimental dataset contains 1,718 Yes examples and 1,933 No examples. Stratified splitting with seed 42 produced 2,555 training examples, 548 validation examples, and 548 test examples. The test split contains 258 Yes and 290 No examples.")
add_table(doc, ["Dataset stage", "Total", "Actionable (Yes)", "Non-Actionable (No)"], [
    ("Source dataset", "3,657", "1,719", "1,938"),
    ("Processed experimental data", "3,651", "1,718", "1,933"),
    ("Held-out test split", "548", "258", "290"),
], [3000, 1500, 2400, 2460])

add_heading(doc, "5.", "Methodology")
add_para(doc, "All three pipelines use the same finalized train, validation, and test files. The training split is used to fit model parameters; the validation split is used for validation and checkpoint selection where implemented; and the test split is reserved for final evaluation. Model outputs are converted to the Actionable class when the predicted probability is at least 0.50.")
add_para(doc, "The evaluation module applies the same positive-class definition and metric calculations to each model. This common evaluation procedure supports a direct comparison of the reported held-out-test results.")

add_heading(doc, "6.", "Preprocessing")
add_para(doc, "The project preprocessing pipeline normalizes text using Unicode NFKC normalization and HTML decoding. It replaces URLs with $LINK, email addresses with $EMAIL, and numeric expressions with $NUM. Remaining HTML tags are removed, and whitespace is normalized.")
add_para(doc, "Records that are empty or meaningless after cleaning are removed. The pipeline also removes identical cleaned text that appears with conflicting labels, because the project does not resolve which label would be correct, and removes exact duplicate cleaned text with the same label. The resulting data are then split stratifiably into 70% training, 15% validation, and 15% test partitions.")

add_heading(doc, "7.", "Model Architectures")
add_heading(doc, "7.1", "TF-IDF + Logistic Regression", level=2)
add_para(doc, "The first pipeline represents cleaned text with a TF-IDF vectorizer and classifies it using Logistic Regression. The vectorizer uses lowercase processing, unigram and bigram features, min_df=2, max_df=0.9, sublinear term-frequency scaling, Unicode accent stripping, and no stop-word list. Logistic Regression uses C=1.0, balanced class weights, the liblinear solver, and a maximum of 2,000 iterations. The TF-IDF vocabulary is fitted only on the training split.")
add_heading(doc, "7.2", "Pretrained Word2Vec + BiLSTM", level=2)
add_para(doc, "The second pipeline constructs a vocabulary from the training text and maps token identifiers to a pretrained Word2Vec embedding matrix. The implementation uses sequences up to 64 tokens, a vocabulary frequency threshold of 2, and a maximum vocabulary size of 20,000. The embedding layer is trainable. A one-layer bidirectional LSTM with hidden dimension 128 processes the embeddings; its forward and backward final states are concatenated and passed through dropout (0.3) and a fully connected two-class output layer.")
add_heading(doc, "7.3", "Pretrained BERT", level=2)
add_para(doc, "The third pipeline fine-tunes bert-base-uncased using AutoTokenizer and AutoModelForSequenceClassification with two output labels. Text is tokenized with truncation and max-length padding to 64 tokens. The implementation retains all BERT parameters as trainable and uses the model's classification output for the binary decision.")

add_heading(doc, "8.", "Training")
add_para(doc, "The TF-IDF plus Logistic Regression model is fitted on the training split. The BiLSTM configuration specifies batches of 32, up to 20 epochs, learning rate 0.001, gradient clipping at 5.0, and patience of 4 validation epochs. Its implementation uses cross-entropy loss and the Adam optimizer, saving the checkpoint with the highest validation F1-score.")
add_para(doc, "The BERT configuration specifies batches of 16, up to 4 epochs, learning rate 2e-5, weight decay 0.01, a linear warmup ratio of 0.1, gradient clipping at 1.0, and patience of 2 validation epochs. It uses cross-entropy loss, AdamW optimization, a linear warmup scheduler, and selects the checkpoint with the highest validation F1-score. In both neural pipelines, the held-out test split is evaluated after training using the selected checkpoint.")

add_heading(doc, "9.", "Evaluation Metrics")
add_para(doc, "Accuracy is the proportion of all test examples classified correctly. Precision measures the proportion of predicted Actionable examples that are actually Actionable. Recall measures the proportion of actual Actionable examples that are detected. The F1-score is the harmonic mean of precision and recall. ROC-AUC summarizes probability-ranking performance across classification thresholds.")
add_para(doc, "The confusion matrix is reported with rows as true labels and columns as predicted labels, in the order No then Yes. In this project, a false negative is an Actionable email-content example predicted as Non-Actionable; this is important because it represents actionable content that could be missed.")

add_heading(doc, "10.", "Results and Comparison")
add_para(doc, "Table 2 presents the final held-out-test results from the repository's comparison artifact. Percentages are displayed to two decimal places. No additional experiments or aggregate statistics are inferred beyond these recorded results.")
add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"], [
    ("TF-IDF + Logistic Regression", "77.74%", "76.98%", "75.19%", "76.08%", "86.07%"),
    ("Pretrained Word2Vec + BiLSTM", "74.82%", "71.74%", "76.74%", "74.16%", "86.54%"),
    ("Pretrained BERT (bert-base-uncased)", "84.85%", "82.05%", "86.82%", "84.37%", "93.16%"),
], [3000, 1272, 1272, 1272, 1272, 1272])
add_table(doc, ["Model", "TP", "TN", "FP", "FN"], [
    ("TF-IDF + Logistic Regression", "194", "232", "58", "64"),
    ("Pretrained Word2Vec + BiLSTM", "198", "212", "78", "60"),
    ("Pretrained BERT (bert-base-uncased)", "224", "241", "49", "34"),
], [4200, 1290, 1290, 1290, 1290])

add_heading(doc, "11.", "Discussion")
add_para(doc, "On the shared held-out test set, BERT records the highest values for accuracy (84.85%), precision (82.05%), recall (86.82%), F1-score (84.37%), and ROC-AUC (93.16%). It also has the fewest false negatives (34) and false positives (49) among the three reported systems. Thus, under this experimental setup, the fine-tuned BERT model provides the strongest recorded overall performance and detects the largest number of Actionable examples (224 of 258).")
add_para(doc, "The TF-IDF plus Logistic Regression model provides a competitive non-neural reference, with 77.74% accuracy and 76.08% F1-score. The Word2Vec plus BiLSTM model has a higher recall (76.74%) and ROC-AUC (86.54%) than the TF-IDF baseline, but its recorded accuracy, precision, and F1-score are lower. These statements are limited to the single shared test split and the configurations recorded in the project repository.")

add_heading(doc, "12.", "Limitations")
for item in (
    "The source records are primarily individual email sentences or snippets, not full structured emails with separate subject and body fields; conclusions therefore concern email-content actionability.",
    "The reported comparison is based on one fixed, stratified split generated with seed 42. The repository does not report repeated runs, cross-validation, or statistical significance testing.",
    "Results reflect the implemented configurations, preprocessing rules, and a maximum input length of 64 tokens for the neural pipelines; they should not be generalized beyond the evaluated dataset without further testing.",
    "The project results measure offline classification performance. They do not establish performance in a deployed email workflow or with other email sources.",
):
    add_bullet(doc, item)

add_heading(doc, "13.", "Conclusion")
add_para(doc, "MailSense evaluates three NLP approaches for binary actionable email-content classification using a common cleaned dataset split and a common held-out test set. The final recorded comparison shows that fine-tuned bert-base-uncased achieves the strongest results across all reported primary metrics: 84.85% accuracy, 82.05% precision, 86.82% recall, 84.37% F1-score, and 93.16% ROC-AUC. The TF-IDF plus Logistic Regression and pretrained Word2Vec plus BiLSTM models provide documented alternative baselines within the same experimental framework. The conclusions are restricted to the Parakweet Labs Ask0729-fixed.txt data and the project procedures recorded in the repository.")

add_heading(doc, "14.", "References")
refs = [
    "[1] Parakweet Labs. Email Intent Dataset, Ask0729-fixed.txt. Available: https://raw.githubusercontent.com/ParakweetLabs/EmailIntentDataSet/master/src/resources/Ask0729-fixed.txt. Accessed for this project dataset.",
    "[2] MailSense project repository. README.md; data_quality_report.json; final_test_comparison.csv; model configuration files; preprocessing, training, and evaluation source files. Internal project artifacts used as the source of truth for this report.",
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(ref)
    set_run_font(r, 10)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())

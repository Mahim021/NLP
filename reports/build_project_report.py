"""Build the MailSense project report (reports/MailSense_Project_Report.docx).

All numbers and figures are generated from the recorded project files in results/.
Run from the project root:  .venv\\Scripts\\python.exe reports\\build_project_report.py
"""

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from PIL import Image  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
FIG_DIR = ROOT / "reports" / "figures"
OUT = ROOT / "reports" / "MailSense_Project_Report.docx"

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
MODEL_LABELS = {
    "Model A: TF-IDF + Logistic Regression": "TF-IDF + Logistic Regression",
    "Model B: LSTM": "Word2Vec + BiLSTM",
    "Model C: BERT": "BERT (bert-base-uncased)",
}


# --------------------------------------------------------------------------- data

comparison = pd.read_csv(RESULTS / "final_test_comparison.csv")
comparison["label"] = comparison["model"].map(MODEL_LABELS)
assert comparison["label"].notna().all(), "unexpected model names in final_test_comparison.csv"
quality = json.loads((RESULTS / "data_quality_report.json").read_text(encoding="utf-8"))
lstm_hist = pd.read_csv(RESULTS / "lstm" / "training_history.csv")
bert_hist = pd.read_csv(RESULTS / "bert" / "training_history.csv")


def pct(value):
    return f"{value * 100:.2f}%"


def diff(a, b):
    """Difference in percentage points between two values as displayed (rounded to 2 decimals)."""
    return f"{round(a * 100, 2) - round(b * 100, 2):.2f}"


def count(value):
    return f"{value:,}"


# ------------------------------------------------------------------------ figures

plt.rcParams.update({
    "font.family": "serif",
    "font.serif": [FONT, "DejaVu Serif"],
    "font.size": 10,
    "axes.edgecolor": "black",
    "axes.linewidth": 0.8,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "xtick.color": "black",
    "ytick.color": "black",
    "savefig.dpi": 300,
    "savefig.bbox": "tight",
    "savefig.facecolor": "white",
})
FIG_DIR.mkdir(parents=True, exist_ok=True)


def comparison_figure(path):
    metrics = [("accuracy", "Accuracy"), ("f1", "F1-score"), ("roc_auc", "ROC-AUC")]
    styles = [
        {"color": "white", "hatch": "////"},
        {"color": "0.65", "hatch": ""},
        {"color": "0.15", "hatch": ""},
    ]
    width = 0.26
    fig, ax = plt.subplots(figsize=(6.2, 3.1))
    for i, (_, row) in enumerate(comparison.iterrows()):
        xs = [m + (i - 1) * width for m in range(len(metrics))]
        values = [row[key] * 100 for key, _ in metrics]
        bars = ax.bar(xs, values, width, label=row["label"], edgecolor="black",
                      linewidth=0.7, **styles[i])
        for bar, value in zip(bars, values):
            ax.text(bar.get_x() + bar.get_width() / 2, value + 1.2, f"{value:.2f}",
                    ha="center", va="bottom", fontsize=7.5)
    ax.set_xticks(range(len(metrics)))
    ax.set_xticklabels([name for _, name in metrics])
    ax.set_ylabel("Score on test set (%)")
    ax.set_ylim(0, 105)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, 1.16), ncol=3, frameon=False, fontsize=8.5)
    fig.savefig(path)
    plt.close(fig)


def loss_figure(history, path):
    best_epoch = int(history.loc[history["val_f1"].idxmax(), "epoch"])
    fig, ax = plt.subplots(figsize=(3.2, 2.5))
    ax.plot(history["epoch"], history["train_loss"], color="black", linestyle="-", marker="o",
            markersize=4, linewidth=1.1, label="Training loss")
    ax.plot(history["epoch"], history["val_loss"], color="black", linestyle="--", marker="s",
            markersize=4, markerfacecolor="white", linewidth=1.1, label="Validation loss")
    ax.axvline(best_epoch, color="0.45", linestyle=":", linewidth=1)
    ax.set_xticks(history["epoch"])
    ax.set_xlabel("Epoch")
    ax.set_ylabel("Cross-entropy loss")
    ax.legend(frameon=False, fontsize=8, loc="lower center", bbox_to_anchor=(0.5, 1.0), ncol=2)
    fig.savefig(path)
    plt.close(fig)


def grayscale_copy(src, dst):
    # The recorded confusion-matrix PNGs use a blue colour map; the report uses grayscale copies.
    Image.open(src).convert("L").save(dst)


comparison_figure(FIG_DIR / "model_comparison.png")
loss_figure(lstm_hist, FIG_DIR / "loss_lstm.png")
loss_figure(bert_hist, FIG_DIR / "loss_bert.png")
for name in ("tfidf_logreg", "lstm", "bert"):
    grayscale_copy(RESULTS / "figures" / f"cm_{name}_test.png", FIG_DIR / f"cm_{name}_test_gray.png")


# ------------------------------------------------------------------- docx helpers

def style_run(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic
    return run


def para(text="", size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_after=6, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = keep_next
    if text:
        add_rich(p, text, size, bold, italic)
    return p


def add_rich(p, text, size=12, bold=False, italic=False):
    """Add text where **segments** are bold."""
    for i, part in enumerate(text.split("**")):
        if part:
            style_run(p.add_run(part), size, bold or i % 2 == 1, italic)


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    add_rich(p, text)
    return p


def heading(number, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    style_run(p.add_run(f"{number} {text}"), 14 if level == 1 else 12, bold=True)
    return p


def caption(text, before=4, after=10):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=after)
    p.paragraph_format.space_before = Pt(before)
    label, _, rest = text.partition(". ")
    style_run(p.add_run(label + ". "), 10, bold=True)
    style_run(p.add_run(rest), 10)
    return p


def set_borders(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        size = edges.get(edge)
        el.set(qn("w:val"), "single" if size else "nil")
        if size:
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), "000000")
        borders.append(el)


def table(headers, rows, widths, caption_text, first_col_left=True):
    """Three-rule academic table with the caption above it."""
    cap = caption(caption_text, before=6, after=4)
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    last = len(rows)
    for r, values in enumerate([headers] + rows):
        row = t.rows[r]
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for c, value in enumerate(values):
            cell = row.cells[c]
            cell.width = Inches(widths[c])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.keep_with_next = r < last
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if c == 0 and first_col_left
                           else WD_ALIGN_PARAGRAPH.CENTER)
            style_run(p.add_run(str(value)), 10.5, bold=(r == 0))
            set_borders(cell, top=12 if r == 0 else None,
                        bottom=12 if r == last else (6 if r == 0 else None))
    para(space_after=4)
    return t


def figure_row(images, width, sub_captions, caption_text):
    """Place one or more images side by side in a borderless table, caption below."""
    t = doc.add_table(rows=2, cols=len(images))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, (img, sub) in enumerate(zip(images, sub_captions)):
        top, bottom = t.rows[0].cells[c], t.rows[1].cells[c]
        for cell in (top, bottom):
            set_borders(cell)
        p = top.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(img), width=Inches(width))
        if sub:
            q = bottom.paragraphs[0]
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.keep_with_next = True
            if sub.startswith("Figure"):
                label, _, rest = sub.partition(". ")
                style_run(q.add_run(label + ". "), 10, bold=True)
                style_run(q.add_run(rest), 10)
            else:
                style_run(q.add_run(sub), 10)
    if caption_text:
        caption(caption_text)


# ------------------------------------------------------------------------ document

doc = Document()
section = doc.sections[0]
for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(section, side, Inches(1))

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
normal.font.color.rgb = BLACK
normal.paragraph_format.line_spacing = 1.15
for name, size, before, after in (("Heading 1", 14, 16, 6), ("Heading 2", 12, 10, 4)):
    s = doc.styles[name]
    s.font.name = FONT
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.italic = False
    s.font.color.rgb = BLACK
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
doc.styles["List Bullet"].font.name = FONT
doc.styles["List Bullet"].font.size = Pt(12)

# Page numbers in the footer of the body section (the title page has none).
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run_el = OxmlElement("w:r")
fld.append(run_el)
footer_p._p.append(fld)
section.different_first_page_header_footer = True

# ---- Title page
for _ in range(6):
    para(space_after=12)
para("MailSense: Actionable Email Detection Using NLP", 20, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
para("Binary Classification of Email Text into Actionable and Non-Actionable", 13,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
para("Project Report", 13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=60)
para("Submitted by", 12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
for line in ("Md Ariful Alam Mahim (Roll: 2107023)", "Md Jubair Husain (Roll: 2107029)"):
    para(line, 12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para("September 2026", 12, align=WD_ALIGN_PARAGRAPH.CENTER).paragraph_format.space_before = Pt(60)
doc.paragraphs[-1].runs[-1].add_break(WD_BREAK.PAGE)

res = {row["label"]: row for _, row in comparison.iterrows()}
A = res["TF-IDF + Logistic Regression"]
B = res["Word2Vec + BiLSTM"]
C = res["BERT (bert-base-uncased)"]
lstm_best = lstm_hist.loc[lstm_hist["val_f1"].idxmax()]
bert_best = bert_hist.loc[bert_hist["val_f1"].idxmax()]
n_test = quality["test_size"]
test_yes = quality["test_class_distribution"]["Yes"]

# ---- 1 Introduction
heading("1", "Introduction")
para("Email inboxes mix content that requires the recipient to do something, such as requests, "
     "proposals and questions, with content that does not, such as newsletters, notifications and "
     "purely informative statements. An overlooked request can lead to a missed deadline or an "
     "unanswered question, so automatically flagging actionable content is a useful aid for email "
     "triage. The task is closely related to earlier work on classifying email into speech acts such "
     "as request and propose [2].")
para("MailSense treats this as a supervised binary text-classification problem and compares three "
     "approaches: (A) TF-IDF features with Logistic Regression, (B) a BiLSTM initialized with pretrained "
     "Word2Vec embeddings, and (C) a fine-tuned pretrained BERT model (bert-base-uncased). All three "
     "models use the same cleaned data, the same train/validation/test split and the same evaluation "
     "code, so that differences in the results reflect the modelling approach rather than differences in "
     "data handling.")

# ---- 2 Problem statement and objectives
heading("2", "Problem Statement and Objectives")
heading("2.1", "Problem Statement", level=2)
para("Given a piece of email text, the system must predict one of two labels: **Actionable (Yes)**, "
     "if the text asks for or proposes an action by the recipient, or **Non-Actionable (No)** otherwise. "
     "Actionable is the positive class. A false negative is therefore an actionable text predicted as "
     "non-actionable, i.e. an item the user could miss; for this reason recall and the number of false "
     "negatives are reported alongside the other metrics. Because the dataset consists of individual "
     "email sentences or short snippets rather than complete messages, the task is more precisely "
     "sentence-level classification of actionable email content.")
heading("2.2", "Objectives", level=2)
for item in (
    "Prepare a cleaned, de-duplicated dataset and a single stratified train/validation/test split "
    "shared by all models.",
    "Implement three classifiers: TF-IDF + Logistic Regression, a BiLSTM initialized with pretrained "
    "Word2Vec embeddings, and fine-tuned BERT.",
    "Evaluate all models on the same held-out test set using accuracy, precision, recall, F1-score, "
    "ROC-AUC and confusion matrices, with Actionable as the positive class.",
    "Compare the results, examine the error types of each model and state the limitations of the study.",
):
    bullet(item)

# ---- 3 Dataset and preprocessing
heading("3", "Dataset and Preprocessing")
heading("3.1", "Dataset Description", level=2)
raw = quality["raw_class_distribution"]
para(f"We use the training file Ask0729-fixed.txt of the Parakweet Labs Email Intent Dataset [1]. Each "
     f"line holds a Yes/No label and one sentence taken from the Enron email corpus [3]; the dataset's "
     f"positive label corresponds primarily to the request and propose speech acts of [2]. The file "
     f"contains **{count(quality['raw_rows'])} examples: {count(raw['Yes'])} Actionable** "
     f"({raw['Yes'] / quality['raw_rows']:.1%}) **and {count(raw['No'])} Non-Actionable** "
     f"({raw['No'] / quality['raw_rows']:.1%}), so the classes are close to balanced. All lines were "
     f"well-formed ({quality['malformed_or_invalid_rows']} rows rejected at parsing).")

heading("3.2", "Data Cleaning and Preparation", level=2)
final = quality["final_class_distribution"]
para("A shared cleaning function is applied to every example: Unicode (NFKC) normalization, decoding of "
     "HTML entities, replacement of URLs, email addresses and numbers with the placeholder tokens $LINK, "
     "$EMAIL and $NUM, removal of HTML tags, and whitespace normalization. Stop words and punctuation are "
     "kept, because words such as please, could and you, and the question mark, are strong cues for "
     "requests. After cleaning, examples that were empty, had conflicting labels, or were exact "
     "duplicates were removed. Only the duplicate filter removed anything: "
     f"{quality['removed_total']} examples, several of which became identical only after email addresses "
     f"were replaced by $EMAIL. The **final dataset contains {count(quality['final_rows'])} examples** "
     f"({count(final['Yes'])} Actionable, {count(final['No'])} Non-Actionable).")

heading("3.3", "Train–Validation–Test Split", level=2)
para(f"The final dataset is divided into 70% training, 15% validation and 15% test data by stratified "
     f"random splitting (seed {quality['seed']}), which preserves the class ratio in each part. The split "
     f"is saved once and reused unchanged by all three models; the test set ({n_test} examples) is used "
     f"only for the final evaluation. Table 1 summarizes the data at each stage.")


def split_row(name, dist, total):
    return [name, count(dist["Yes"]), count(dist["No"]), count(total)]


table(["Stage", "Actionable (Yes)", "Non-Actionable (No)", "Total"], [
    split_row("Raw dataset", raw, quality["raw_rows"]),
    split_row("After cleaning", final, quality["final_rows"]),
    split_row("Training set (70%)", quality["train_class_distribution"], quality["train_size"]),
    split_row("Validation set (15%)", quality["val_class_distribution"], quality["val_size"]),
    split_row("Test set (15%)", quality["test_class_distribution"], n_test),
], [2.1, 1.5, 1.8, 0.9], "Table 1. Dataset size and class distribution at each stage.")

# ---- 4 Methodology
heading("4", "Methodology")
heading("4.1", "TF-IDF + Logistic Regression", level=2)
para("The baseline represents each text as a sparse TF-IDF vector [4] of lower-cased word unigrams and "
     "bigrams, with sub-linear term-frequency scaling and n-grams kept only if they occur in at least 2 "
     "and at most 90% of training documents. The vectorizer is fitted on the training split only, giving "
     "6,572 features. A Logistic Regression classifier with L2 regularization (C = 1.0), balanced class "
     "weights and the liblinear solver (maximum 2,000 iterations) is trained on these vectors [9].")

heading("4.2", "Word2Vec + BiLSTM", level=2)
para("Model B is a BiLSTM initialized with pretrained Word2Vec embeddings. A vocabulary of 2,450 "
     "entries (words with frequency of at least 2 in the training split, plus padding and unknown "
     "tokens) is built, and its 300-dimensional embedding matrix is initialized from the pretrained Google "
     "News Word2Vec vectors [5]; 2,338 of 2,448 vocabulary words (95.51%) were found, and the remaining "
     "words were initialized randomly. The embeddings are trainable and are fine-tuned together with the "
     "rest of the network. A single-layer bidirectional LSTM [6, 7] with 128 hidden units per direction "
     "reads the sequence (maximum 64 tokens); the final forward and backward states are concatenated and "
     "passed through dropout (p = 0.3) and a linear output layer. The model has 1,175,834 trainable "
     "parameters.")

heading("4.3", "Pretrained BERT", level=2)
para("Model C fine-tunes bert-base-uncased [8], a 12-layer Transformer encoder that produces contextual "
     "token representations. It is loaded through the Hugging Face Transformers library [11] with a "
     "two-class sequence-classification head (a linear layer on the [CLS] representation). Inputs are "
     "WordPiece-tokenized and truncated or padded to 64 tokens, and all 109,483,778 parameters are "
     "fine-tuned.")

heading("4.4", "Experimental Setup", level=2)
para("The experiments were run in a single Kaggle notebook on an NVIDIA Tesla T4 GPU using scikit-learn "
     "[9], PyTorch [10], Transformers [11] and gensim [12]. The random seed is fixed to 42 for all models. "
     "The TF-IDF model is fitted once on the training set with the fixed configuration above. The two "
     "neural models minimize cross-entropy loss; after every epoch they are evaluated on the validation "
     "set, the checkpoint with the highest validation F1-score is kept, and training stops early if the "
     "validation F1 does not improve for a set number of epochs (patience). The selected checkpoint is "
     "then evaluated once on the test set. No hyperparameter search was performed; the settings in "
     "Table 2 were used as fixed.")
table(["Setting", "Word2Vec + BiLSTM", "BERT"], [
    ["Optimizer", "Adam [14]", "AdamW [13]"],
    ["Learning rate", "1 × 10⁻³, constant", "2 × 10⁻⁵, linear decay, 10% warm-up"],
    ["Weight decay", "0", "0.01"],
    ["Batch size", "32", "16"],
    ["Maximum epochs / patience", "20 / 4", "4 / 2"],
    ["Gradient-norm clipping", "5.0", "1.0"],
    ["Maximum sequence length", "64 tokens", "64 WordPiece tokens"],
], [2.2, 1.9, 2.3], "Table 2. Training settings of the neural models.")

# ---- 5 Evaluation metrics
heading("5", "Evaluation Metrics")
para(f"All models are evaluated by the same code on the same test set of {n_test} examples "
     f"({test_yes} Actionable, {quality['test_class_distribution']['No']} Non-Actionable). A text is "
     "predicted Actionable when the model's predicted probability of the Actionable class is at least "
     "**0.5**. With Actionable as the positive class, TP, TN, FP and FN denote true positives, true "
     "negatives, false positives and false negatives.")
for item in (
    "**Accuracy** = (TP + TN) / (TP + TN + FP + FN): the proportion of all texts classified correctly.",
    "**Precision** = TP / (TP + FP): the proportion of texts flagged as actionable that are actionable.",
    "**Recall** = TP / (TP + FN): the proportion of actionable texts that are found.",
    "**F1-score** = 2 · Precision · Recall / (Precision + Recall): the harmonic mean of "
    "precision and recall; also the model-selection criterion on the validation set.",
    "**ROC-AUC**: the area under the ROC curve, computed from the predicted probabilities. It measures "
    "how well actionable texts are ranked above non-actionable ones across all thresholds, independently "
    "of the 0.5 threshold.",
    "**Confusion matrix**: the counts of TP, TN, FP and FN, with true labels as rows and predicted labels "
    "as columns.",
):
    bullet(item)

# ---- 6 Results
heading("6", "Results and Analysis")
heading("6.1", "Overall Model Performance", level=2)
para(f"Table 3 reports the test-set results of the single final run, taken from the recorded result files. "
     "Figure 1 compares accuracy, F1-score and ROC-AUC.")
table(["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"], [
    [row["label"], pct(row["accuracy"]), pct(row["precision"]), pct(row["recall"]), pct(row["f1"]),
     pct(row["roc_auc"])] for _, row in comparison.iterrows()
], [2.3, 0.85, 0.85, 0.8, 0.8, 0.85],
    f"Table 3. Test-set results (n = {n_test}; positive class = Actionable; threshold = 0.5).")
figure_row([FIG_DIR / "model_comparison.png"], 5.8, [None],
           "Figure 1. Accuracy, F1-score and ROC-AUC of the three models on the test set.")
para(f"BERT achieved the highest measured value on all five metrics among the three models in this "
     f"experiment, exceeding the TF-IDF baseline by {diff(C['accuracy'], A['accuracy'])} "
     f"percentage points in accuracy, {diff(C['f1'], A['f1'])} points in F1-score and "
     f"{diff(C['roc_auc'], A['roc_auc'])} points in ROC-AUC. The TF-IDF baseline scored higher "
     f"than the BiLSTM on accuracy, precision and F1-score, whereas the BiLSTM had higher recall "
     f"({pct(B['recall'])} vs. {pct(A['recall'])}) and a slightly higher ROC-AUC "
     f"({pct(B['roc_auc'])} vs. {pct(A['roc_auc'])}).")

heading("6.2", "Confusion Matrix Analysis", level=2)
para(f"Figure 2 shows the test-set confusion matrices. Of the {test_yes} actionable test examples, BERT "
     f"missed {C['FN']} (false negatives), compared with {A['FN']} for TF-IDF + Logistic Regression and "
     f"{B['FN']} for the BiLSTM. BERT also produced the fewest false positives ({C['FP']}). The BiLSTM "
     f"produced the most false positives ({B['FP']}, against {A['FP']} for the baseline), which explains "
     "its lower precision despite its slightly higher recall.")
figure_row([FIG_DIR / f"cm_{m}_test_gray.png" for m in ("tfidf_logreg", "lstm", "bert")], 2.05,
           ["(a) TF-IDF + Logistic Regression", "(b) Word2Vec + BiLSTM", "(c) BERT"],
           "Figure 2. Test-set confusion matrices (rows: true label; columns: predicted label).")

heading("6.3", "Training Performance", level=2)
para(f"Figures 3 and 4 plot the training and validation loss recorded for the two neural models. For the "
     f"BiLSTM, the best validation F1-score ({pct(lstm_best['val_f1'])}) was reached at epoch "
     f"{int(lstm_best['epoch'])}; it did not improve over the next four epochs, so training stopped after "
     f"epoch {int(lstm_hist['epoch'].max())} and the epoch-{int(lstm_best['epoch'])} checkpoint was used. "
     f"The training loss kept falling (to {lstm_hist['train_loss'].iloc[-1]:.3f}) while the validation "
     f"loss rose after epoch {int(lstm_hist.loc[lstm_hist['val_loss'].idxmin(), 'epoch'])} (to "
     f"{lstm_hist['val_loss'].iloc[-1]:.3f}), indicating overfitting.")
para(f"For BERT, the validation F1-score improved in every epoch, so all {int(bert_hist['epoch'].max())} "
     f"epochs were run and the final checkpoint was selected (validation F1 {pct(bert_best['val_f1'])}). "
     f"Its validation loss was lowest after epoch "
     f"{int(bert_hist.loc[bert_hist['val_loss'].idxmin(), 'epoch'])} and then increased although "
     "validation accuracy and F1 improved, which suggests the model became more confident on the examples "
     "it still misclassified. Checkpoints were selected by validation F1, not by validation loss.")
figure_row([FIG_DIR / "loss_lstm.png", FIG_DIR / "loss_bert.png"], 3.0,
           ["Figure 3. BiLSTM training loss vs. validation loss.",
            "Figure 4. BERT training loss vs. validation loss."],
           None)
para("The dotted vertical line marks the checkpoint selected by validation F1-score.", 10,
     italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10).paragraph_format.space_before = Pt(4)

heading("6.4", "Discussion", level=2)
para("**Contextual pretraining.** Whether a sentence is actionable often depends on its form and on how "
     "words combine, not only on individual words: the same verb appears in \"Could you send the file?\" "
     "and \"I sent the file.\" BERT represents each token in the context of the whole sentence and brings "
     "knowledge from large-scale pretraining, which plausibly explains its higher scores here. This comes "
     "at a much higher cost: about 109.5 million fine-tuned parameters and roughly 30 seconds per epoch on "
     "the T4 GPU, against about one second per epoch for the BiLSTM.")
para("**A strong lexical baseline.** Actionable sentences often contain explicit cues such as please, "
     "let me know or can you, which unigram and bigram features capture directly. With only 2,555 "
     "training examples, a regularized linear model is less prone to overfitting than a network with "
     "over a million parameters, which is consistent with TF-IDF + Logistic Regression outperforming the "
     "BiLSTM on test accuracy and F1-score.")
para(f"**Generalization of the BiLSTM.** The BiLSTM reached a validation F1-score of "
     f"{pct(lstm_best['val_f1'])} but a test F1-score of {pct(B['f1'])}, the largest drop of the three "
     "models. Together with the rising validation loss, this suggests that its selected checkpoint fitted "
     "the particular validation sample and generalized less well. Its errors were skewed towards false "
     "positives. Its ROC-AUC was nevertheless slightly above the baseline's, so the two models rank "
     "examples comparably well overall and part of the difference in thresholded metrics depends on how "
     "their probabilities fall relative to the fixed 0.5 threshold.")

# ---- 7 Limitations
heading("7", "Limitations")
for item in (
    "**Single train/validation/test split.** All results come from one stratified split; performance on "
    "other splits was not measured, and no cross-validation was performed.",
    f"**Limited dataset size.** The final dataset has {count(quality['final_rows'])} examples and the test "
    f"set {n_test}, so small differences, such as the "
    f"{diff(B['roc_auc'], A['roc_auc'])}-point ROC-AUC difference between the BiLSTM and the "
    "baseline, should not be treated as conclusive.",
    "**Sentence-level data.** The examples are individual sentences or snippets, not complete emails; "
    "the models do not use the rest of the message, the subject line, the sender or the thread.",
    "**Fixed threshold.** All models used a decision threshold of 0.5. A threshold chosen to favour "
    "recall could reduce missed actionable items at the cost of more false positives; this was not "
    "explored.",
    "**Single experimental run.** Each model was trained once with one random seed and one fixed "
    "configuration, without a hyperparameter search. Variability across seeds was not measured and no "
    "statistical significance tests were performed.",
):
    bullet(item)

# ---- 8 Conclusion
heading("8", "Conclusion")
para(f"MailSense compared three NLP approaches for detecting actionable email content under a controlled "
     f"protocol: {count(quality['final_rows'])} cleaned examples from the Parakweet Labs Email Intent "
     f"Dataset, one stratified 70/15/15 split shared by all models, and the same evaluation code. On the "
     f"{n_test}-example test set, fine-tuned BERT achieved the highest measured performance among the three "
     f"models in this experiment, with {pct(C['accuracy'])} accuracy, {pct(C['f1'])} F1-score and "
     f"{pct(C['roc_auc'])} ROC-AUC, and it missed the fewest actionable examples ({C['FN']} of {test_yes}). "
     f"TF-IDF + Logistic Regression ({pct(A['accuracy'])} accuracy, {pct(A['f1'])} F1) remained a "
     f"competitive low-cost baseline, while the BiLSTM initialized with pretrained Word2Vec embeddings "
     f"({pct(B['accuracy'])} accuracy, {pct(B['f1'])} F1) showed higher recall but lower precision than "
     "the baseline and signs of overfitting.")
para("Natural next steps, following from the limitations, are to repeat the experiments over several "
     "seeds and splits with significance testing, to tune the decision threshold for higher recall, and "
     "to extend the task from single sentences to complete email messages.")

# ---- References
heading("", "References").runs[0].text = "References"
refs = [
    "Parakweet Labs, Inc., “EmailIntentDataSet: Labeled training and test data for email intent "
    "machine learning (for sentence-level speech acts),” GitHub repository, 2014. "
    "https://github.com/ParakweetLabs/EmailIntentDataSet",
    "W. W. Cohen, V. R. Carvalho, and T. M. Mitchell, “Learning to classify email into ‘speech "
    "acts’,” in Proc. EMNLP, 2004, pp. 309–316.",
    "B. Klimt and Y. Yang, “The Enron corpus: A new dataset for email classification research,” "
    "in Proc. ECML, 2004, pp. 217–226.",
    "G. Salton and C. Buckley, “Term-weighting approaches in automatic text retrieval,” "
    "Information Processing & Management, vol. 24, no. 5, pp. 513–523, 1988.",
    "T. Mikolov, I. Sutskever, K. Chen, G. S. Corrado, and J. Dean, “Distributed representations of "
    "words and phrases and their compositionality,” in Advances in NIPS, vol. 26, 2013.",
    "S. Hochreiter and J. Schmidhuber, “Long short-term memory,” Neural Computation, vol. 9, "
    "no. 8, pp. 1735–1780, 1997.",
    "M. Schuster and K. K. Paliwal, “Bidirectional recurrent neural networks,” IEEE Trans. "
    "Signal Processing, vol. 45, no. 11, pp. 2673–2681, 1997.",
    "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, “BERT: Pre-training of deep bidirectional "
    "transformers for language understanding,” in Proc. NAACL-HLT, 2019, pp. 4171–4186.",
    "F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” JMLR, vol. 12, "
    "pp. 2825–2830, 2011.",
    "A. Paszke et al., “PyTorch: An imperative style, high-performance deep learning library,” "
    "in Advances in NeurIPS, vol. 32, 2019.",
    "T. Wolf et al., “Transformers: State-of-the-art natural language processing,” in Proc. "
    "EMNLP 2020: System Demonstrations, pp. 38–45.",
    "R. Řehůřek and P. Sojka, “Software framework for topic modelling with large "
    "corpora,” in Proc. LREC 2010 Workshop on New Challenges for NLP Frameworks, pp. 45–50.",
    "I. Loshchilov and F. Hutter, “Decoupled weight decay regularization,” in Proc. ICLR, 2019.",
    "D. P. Kingma and J. Ba, “Adam: A method for stochastic optimization,” in Proc. ICLR, 2015.",
]
for i, ref in enumerate(refs, 1):
    p = para(align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    style_run(p.add_run(f"[{i}]\t{ref}"), 10.5)

doc.save(OUT)
print(OUT)

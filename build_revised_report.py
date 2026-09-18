from pathlib import Path
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "reports" / "MailSense_Academic_Project_Report_Revised.docx"
FIG = ROOT / "reports" / "figures"
BLACK = RGBColor(0, 0, 0)


def graph_style():
    plt.rcParams.update({"font.family": "serif", "font.size": 10, "axes.edgecolor": "black", "axes.labelcolor": "black", "xtick.color": "black", "ytick.color": "black"})


def make_loss_plot(csv_name, title, outfile):
    df = pd.read_csv(ROOT / csv_name)
    graph_style()
    fig, ax = plt.subplots(figsize=(6.2, 3.4))
    ax.plot(df["epoch"], df["train_loss"], color="black", marker="o", linewidth=1.5, label="Training loss")
    ax.plot(df["epoch"], df["val_loss"], color="0.45", marker="s", linewidth=1.5, label="Validation loss")
    ax.set_title(title, fontsize=11, fontweight="bold")
    ax.set_xlabel("Epoch")
    ax.set_ylabel("Cross-entropy loss")
    ax.set_xticks(df["epoch"])
    ax.grid(axis="y", color="0.85", linewidth=0.6)
    ax.legend(frameon=False, loc="best")
    fig.tight_layout()
    fig.savefig(outfile, dpi=200, facecolor="white")
    plt.close(fig)


def make_comparison_plot(outfile):
    models = ["TF-IDF + LR", "Word2Vec +\nBiLSTM", "BERT"]
    metrics = {
        "Accuracy": [77.74, 74.82, 84.85],
        "F1": [76.08, 74.16, 84.37],
        "ROC-AUC": [86.07, 86.54, 93.16],
    }
    graph_style()
    fig, ax = plt.subplots(figsize=(6.2, 3.6))
    styles = [("black", "o"), ("0.40", "s"), ("0.70", "^")]
    for (name, values), (color, marker) in zip(metrics.items(), styles):
        ax.plot(models, values, color=color, marker=marker, linewidth=1.5, label=name)
    ax.set_ylim(65, 100)
    ax.set_ylabel("Score (%)")
    ax.set_title("Held-out Test Performance", fontsize=11, fontweight="bold")
    ax.grid(axis="y", color="0.85", linewidth=0.6)
    ax.legend(frameon=False, ncol=3, loc="lower center")
    fig.tight_layout()
    fig.savefig(outfile, dpi=200, facecolor="white")
    plt.close(fig)


def make_grayscale_confusion_matrices():
    source = ROOT / "results" / "figures"
    for stem in ("cm_tfidf_logreg_test", "cm_lstm_test", "cm_bert_test"):
        with Image.open(source / f"{stem}.png") as image:
            image.convert("L").save(FIG / f"{stem}_grayscale.png")


def set_run(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "000000")


def setup_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell._tc.tcPr.tcW.set(qn("w:w"), str(int(width * 1440)))
            cell._tc.tcPr.tcW.set(qn("w:type"), "dxa")
            set_cell_border(cell)


def add_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    setup_table(t, widths)
    for cell, value in zip(t.rows[0].cells, headers):
        cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd"))
        cell._tc.tcPr[-1].set(qn("w:fill"), "E6E6E6")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(value), 9.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if isinstance(value, str) and value.endswith("%") else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(value), 9.5)
    return t


def add_p(doc, text, size=11, align=None, bold=False, italic=False, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run(r, size, bold, italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), 14 if level == 1 else 12, bold=True)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run(text), 9, italic=True)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(paragraph.add_run("Page "), 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


FIG.mkdir(parents=True, exist_ok=True)
make_loss_plot("results/lstm/training_history.csv", "BiLSTM Training and Validation Loss", FIG / "lstm_loss.png")
make_loss_plot("results/bert/training_history.csv", "BERT Training and Validation Loss", FIG / "bert_loss.png")
make_comparison_plot(FIG / "model_comparison.png")
make_grayscale_confusion_matrices()

doc = Document()
sec = doc.sections[0]
for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, side, Inches(1))
sec.header_distance = Inches(0.49)
sec.footer_distance = Inches(0.49)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(11)
normal.font.color.rgb = BLACK
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(6)
for style, size, before, after in (("Heading 1", 14, 12, 5), ("Heading 2", 12, 9, 3)):
    s = doc.styles[style]
    s.font.name = "Times New Roman"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = BLACK
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
footer = sec.footer.paragraphs[0]
page_number(footer)

# Plain title page.
for _ in range(7): doc.add_paragraph()
add_p(doc, "MailSense: Actionable Email Detection Using NLP", 19, WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=16)
add_p(doc, "Academic Project Report", 13, WD_ALIGN_PARAGRAPH.CENTER, after=36)
add_p(doc, "Md Ariful Alam Mahim (Roll: 2107023)", 12, WD_ALIGN_PARAGRAPH.CENTER, after=5)
add_p(doc, "Md Jubair Husain (Roll: 2107029)", 12, WD_ALIGN_PARAGRAPH.CENTER, after=32)
add_p(doc, "September 2026", 11, WD_ALIGN_PARAGRAPH.CENTER)
doc.add_section(WD_SECTION.NEW_PAGE)

heading(doc, "1. Introduction")
add_p(doc, "MailSense is a binary NLP classifier that identifies whether an email-text snippet is Actionable (Yes) or Non-Actionable (No). The project compares a TF-IDF plus Logistic Regression baseline, a BiLSTM initialized with pretrained Word2Vec embeddings, and fine-tuned bert-base-uncased. Because the data are email sentences or snippets rather than complete subject-and-body messages, the scope is actionable email-content classification.")

heading(doc, "2. Problem Statement and Objectives")
heading(doc, "2.1 Problem Statement", 2)
add_p(doc, "Given an email-text snippet, predict Yes (Actionable) or No (Non-Actionable). The experiment evaluates the three implemented models on the same held-out test set.")
heading(doc, "2.2 Objectives", 2)
for item in ("Prepare a common cleaned dataset split for all models.", "Train and evaluate the three specified classification pipelines.", "Compare held-out-test accuracy, precision, recall, F1-score, ROC-AUC, and confusion matrices."):
    p = doc.add_paragraph(style="List Bullet")
    set_run(p.add_run(item))

heading(doc, "3. Dataset and Preprocessing")
heading(doc, "3.1 Dataset Description", 2)
add_p(doc, "The Parakweet Labs Email Intent Dataset file Ask0729-fixed.txt contains 3,657 raw examples: 1,719 Actionable (Yes) and 1,938 Non-Actionable (No). These are source-data counts, not post-processing experimental counts.")
heading(doc, "3.2 Data Cleaning and Preparation", 2)
add_p(doc, "The implemented pipeline applies Unicode normalization and HTML decoding; replaces URLs, email addresses, and numbers with $LINK, $EMAIL, and $NUM; removes remaining HTML tags; normalizes whitespace; and removes meaningless, conflicting-label, and duplicate cleaned records. Six records were removed, leaving 3,651 examples (1,718 Yes; 1,933 No) for the experiment.")
heading(doc, "3.3 Train-Validation-Test Split", 2)
add_p(doc, "A stratified 70/15/15 split with seed 42 produced 2,555 training, 548 validation, and 548 test examples. The same split was used for every model. The held-out test set contains 258 Actionable and 290 Non-Actionable examples.")

heading(doc, "4. Methodology")
heading(doc, "4.1 TF-IDF + Logistic Regression", 2)
add_p(doc, "Text was represented with TF-IDF unigram and bigram features (min_df=2, max_df=0.9, sublinear term frequency) and classified with balanced Logistic Regression (C=1.0). The vectorizer was fitted on training data only.")
heading(doc, "4.2 Word2Vec + BiLSTM", 2)
add_p(doc, "This model is a BiLSTM initialized with pretrained Word2Vec embeddings. The embeddings were trainable during training. The implementation uses a maximum length of 64 tokens, a one-layer bidirectional LSTM with hidden dimension 128, dropout of 0.3, and a two-class output layer.")
heading(doc, "4.3 Pretrained BERT", 2)
add_p(doc, "bert-base-uncased was fine-tuned for two-class sequence classification. Inputs were tokenized and padded or truncated to 64 tokens; all BERT parameters remained trainable.")
heading(doc, "4.4 Experimental Setup", 2)
add_p(doc, "The BiLSTM configuration used batches of 32, up to 20 epochs, Adam optimization, and validation-F1 checkpoint selection. BERT used batches of 16, up to 4 epochs, AdamW with linear warmup, and validation-F1 checkpoint selection. For all models, Actionable (Yes) was the positive class and a probability threshold of 0.5 produced the final class prediction.")

heading(doc, "5. Evaluation Metrics")
add_p(doc, "Accuracy, precision, recall, F1-score, and ROC-AUC were computed on the held-out test set. Precision, recall, and F1-score use Actionable (Yes) as the positive class. Confusion matrices report true labels by rows and predicted labels by columns; a false negative is actionable content predicted as non-actionable.")

heading(doc, "6. Results and Analysis")
heading(doc, "6.1 Overall Model Performance", 2)
add_p(doc, "Table 1 reports the final held-out-test results recorded in the project files.")
add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"], [
    ("TF-IDF + Logistic Regression", "77.74%", "76.98%", "75.19%", "76.08%", "86.07%"),
    ("Word2Vec + BiLSTM", "74.82%", "71.74%", "76.74%", "74.16%", "86.54%"),
    ("BERT (bert-base-uncased)", "84.85%", "82.05%", "86.82%", "84.37%", "93.16%"),
], [2.35, 0.83, 0.83, 0.83, 0.83, 0.83])
caption(doc, "Table 1. Final metrics on the common held-out test set (n=548).")
doc.add_picture(str(FIG / "model_comparison.png"), width=Inches(5.9))
caption(doc, "Figure 1. Accuracy, F1, and ROC-AUC across the three models.")

heading(doc, "6.2 Confusion Matrix Analysis", 2)
add_p(doc, "The recorded confusion matrices show that BERT produced 224 true positives, 241 true negatives, 49 false positives, and 34 false negatives. TF-IDF plus Logistic Regression recorded 194, 232, 58, and 64 respectively; Word2Vec plus BiLSTM recorded 198, 212, 78, and 60 respectively.")
for path, label in (("cm_tfidf_logreg_test_grayscale.png", "Figure 2. TF-IDF + Logistic Regression test confusion matrix."), ("cm_lstm_test_grayscale.png", "Figure 3. Word2Vec + BiLSTM test confusion matrix."), ("cm_bert_test_grayscale.png", "Figure 4. BERT test confusion matrix.")):
    doc.add_picture(str(FIG / path), width=Inches(4.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, label)

heading(doc, "6.3 Training Performance", 2)
add_p(doc, "The following plots reproduce the training and validation loss values from the recorded training-history CSV files; no values were added or estimated.")
doc.add_picture(str(FIG / "lstm_loss.png"), width=Inches(5.9))
caption(doc, "Figure 5. BiLSTM training and validation loss by epoch.")
doc.add_picture(str(FIG / "bert_loss.png"), width=Inches(5.9))
caption(doc, "Figure 6. BERT training and validation loss by epoch.")

heading(doc, "6.4 Discussion", 2)
add_p(doc, "Among the three models in this experiment, BERT achieved the highest measured accuracy (84.85%), precision (82.05%), recall (86.82%), F1-score (84.37%), and ROC-AUC (93.16%). It also had the fewest recorded false negatives (34), whereas the TF-IDF and BiLSTM models had 64 and 60. The TF-IDF baseline achieved 77.74% accuracy and 76.08% F1. The BiLSTM recorded higher recall and ROC-AUC than the TF-IDF baseline, but lower accuracy, precision, and F1 on this test split.")

heading(doc, "7. Limitations")
for item in ("The results come from one stratified train/validation/test split and a single experimental run.", "The dataset is limited in size and consists primarily of sentences or snippets rather than complete emails.", "All final class decisions use the fixed 0.5 probability threshold.", "The recorded findings apply to the evaluated dataset and implemented configurations; no additional datasets, searches, or experiments are reported."):
    p = doc.add_paragraph(style="List Bullet")
    set_run(p.add_run(item))

heading(doc, "8. Conclusion")
add_p(doc, "MailSense compares three NLP pipelines for actionable email-content classification using one common cleaned split. On the held-out test set of 548 examples, fine-tuned BERT achieved the highest measured performance among the three models, with 84.85% accuracy, 84.37% F1-score, and 93.16% ROC-AUC. The results are bounded by the dataset, fixed split, and configurations documented in this project.")

heading(doc, "References")
for reference in (
    "[1] Parakweet Labs, Email Intent Dataset, Ask0729-fixed.txt. Available: https://raw.githubusercontent.com/ParakweetLabs/EmailIntentDataSet/master/src/resources/Ask0729-fixed.txt.",
    "[2] MailSense project artifacts: README.md; results/data_quality_report.json; results/final_test_comparison.csv; results/*/training_history.csv; configs/*.json; and source preprocessing, training, and evaluation files.",
):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_run(p.add_run(reference), 10)

doc.save(OUT)
print(OUT)
